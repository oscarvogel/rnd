from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

base = Path(r'C:\Programacion\RND\docs\manual_hoja_ruta')
shots = base / 'capturas'
out = Path(r'C:\Programacion\RND\docs\manual_usuario_hoja_ruta.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27); sec.page_height = Inches(11.69)
sec.top_margin = Inches(.62); sec.bottom_margin = Inches(.62)
sec.left_margin = Inches(.68); sec.right_margin = Inches(.68)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Aptos')
styles['Normal'].font.size = Pt(10); styles['Normal'].font.color.rgb = RGBColor(31, 41, 55)
styles['Normal'].paragraph_format.space_after = Pt(6); styles['Normal'].paragraph_format.line_spacing = 1.1
for name, size, color in [('Title', 28, '0B3975'), ('Heading 1', 19, '0B3975'), ('Heading 2', 13, '0A84D8')]:
    st=styles[name]; st.font.name='Aptos Display'; st._element.rPr.rFonts.set(qn('w:eastAsia'),'Aptos Display'); st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(9); st.paragraph_format.space_after=Pt(7)
styles['Caption'].font.name='Aptos'; styles['Caption'].font.size=Pt(8); styles['Caption'].font.italic=True; styles['Caption'].font.color.rgb=RGBColor(71,85,105)

# Footer with page number field.
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('RND · Circuito de reparto     |     Página '); run.font.size=Pt(8); run.font.color.rgb=RGBColor(100,116,139)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)

# Page 1: overview
p=doc.add_paragraph(style='Title'); p.add_run('Manual de usuario')
p=doc.add_paragraph(); r=p.add_run('Importar pedidos e imprimir la hoja de ruta'); r.bold=True; r.font.size=Pt(17); r.font.color.rgb=RGBColor(10,132,216)
p=doc.add_paragraph(); r=p.add_run('RND · Circuito de reparto · 20 de octubre de 2026'); r.font.size=Pt(9); r.font.color.rgb=RGBColor(100,116,139)
doc.add_paragraph('Este manual explica el recorrido operativo desde la pantalla principal de RND hasta la impresión de la hoja de ruta: importar el archivo del proveedor, revisar y grabar los pedidos, organizarlos por ruta, asignar chofer y camión, preparar la carga en pallets y generar el PDF.')
p=doc.add_paragraph(); r=p.add_run('Las capturas muestran pantallas actuales de RND. Los clientes, comprobantes y cantidades que aparecen son datos de demostración.'); r.italic=True; r.font.color.rgb=RGBColor(71,85,105)
doc.add_heading('Recorrido rápido', level=1)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Importar pedidos  →  Revisar vista previa  →  Grabar pedidos\nOrganizar pedidos  →  Asignar chofer y camión  →  Armar pallets  →  Validar carga\nMarcar LISTA  →  Confirmar salida a reparto  →  Revisar e imprimir'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(11,57,117)
doc.add_heading('Antes de empezar', level=1)
for t in ['Tenga el archivo recibido del proveedor y confirme la fecha de reparto.','Verifique que el proveedor, los clientes, los lugares de entrega, las rutas, los choferes y los camiones estén cargados en RND.','Si una opción no aparece en la pantalla principal, puede depender del perfil de acceso.']:
    doc.add_paragraph(t, style='List Bullet')

# Reusable image and paragraph functions.
def step(title, paras, image, caption, width=6.9):
    doc.add_page_break(); doc.add_heading(title, level=1)
    for text in paras: doc.add_paragraph(text)
    pic=doc.add_paragraph(); pic.alignment=WD_ALIGN_PARAGRAPH.CENTER; pic.paragraph_format.space_before=Pt(3); pic.paragraph_format.space_after=Pt(2)
    pic.add_run().add_picture(str(shots/image), width=Inches(width))
    cap=doc.add_paragraph(caption, style='Caption'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER

step('1. Importar y revisar el archivo', [
    'Desde la pantalla principal, abra Importar pedidos. Seleccione el Proveedor / origen y la Fecha reparto antes de elegir el archivo. Use Seleccionar archivo para buscarlo en el equipo.',
    'RND admite planillas Excel (.xlsx y .xls), archivos PDF e imágenes (.png, .jpg y .jpeg). En Excel, seleccione la hoja si corresponde; las filas de inicio y fin son opcionales. Los PDF e imágenes se procesan con IA. No hace falta convertirlos manualmente a Excel.',
    'Presione Cargar vista previa y controle cliente, comprobante, producto, cantidad, KG, bultos y observaciones. La vista previa todavía no graba pedidos. En Excel, use la casilla Importa para incluir o excluir filas. En PDF o imagen, corrija los datos extraídos y revise las filas marcadas REVISAR IA.',
    'Cuando la vista previa esté correcta, presione Grabar pedidos. Lea el resumen de importación y continúe con el botón que RND indique. Si hay pendientes, resuélvalos antes de seguir.'
], '01_importar_vista_previa.png', 'Figura 1. Importación con vista previa. Los datos visibles son de demostración.')

step('2. Organizar los pedidos en una ruta', [
    'Abra Pedidos para organizar desde el paso siguiente o desde el menú operativo. Seleccione la fecha y presione Actualizar. Cada fila representa una factura o documento completo; revise el estado, el cliente, el lugar de entrega y la ruta.',
    'Use Editar para corregir cliente, lugar, ruta, remito u observación general. Use Productos para revisar las líneas y, si hace falta, cambiar únicamente la Cantidad a entregar. No puede ser negativa ni superar la cantidad original.',
    'Marque las facturas que viajarán juntas, seleccione la Ruta destino y presione Organizar seleccionados. Espere la confirmación. Después de organizar correctamente, RND habilita Asignar chofer y camión.'
], '02_pedidos_para_organizar.png', 'Figura 2. Bandeja de pedidos agrupada por factura. Los datos visibles son de demostración.')

step('3. Asignar chofer y camión', [
    'En Asignar chofer y camión, seleccione la fecha y la ruta y presione Cargar datos de hoja. Compruebe la cantidad de pedidos, el peso y los bultos antes de asignar recursos.',
    'Seleccione un Chofer / responsable y un Camión / equipo reales. Presione Guardar asignación. RND guarda los recursos para todos los pedidos de esa fecha y ruta. Luego puede presionar Revisar hoja de ruta.'
], '03_asignar_chofer_camion.png', 'Figura 3. Asignación de recursos para una fecha y ruta. Los datos visibles son de demostración.')

step('4. Armar pallets y validar la carga', [
    'El circuito admite dos formas de preparar la mercadería. Si no arman pallets, el circuito se completa igual: el requisito de carga del checklist queda como informativo y no bloquea. Si arman pallets, este paso es el que libera el checklist.',
    'Abra Armado de pallets y seleccione la Fecha de reparto y la Ruta de reparto. Presione Cargar mercadería: a la izquierda aparece la Mercadería pendiente de preparar, con el saldo de cada línea ya descontado de lo asignado a otros pallets.',
    'Seleccione una o más líneas y presione Agregar selección al pallet. Use Agregar parcial… para cargar una cantidad menor al saldo y Quitar del pallet para sacar líneas. Con + Nuevo pallet abre otro pallet; los pallets se muestran numerados, un clic cambia el pallet activo y el KPI muestra su peso. Etiqueta / QR imprime la etiqueta del pallet seleccionado.',
    'Cuando el pallet esté completo, presione Confirmar preparación. Ojo con este punto: Confirmar preparación no cambia el estado. El pallet sigue ARMADO y lo que se ve es el resumen. Para despachar falta el paso siguiente.'
], '04_armado_pallets.png', 'Figura 4. Armado de pallets: mercadería pendiente de preparar y contenido del pallet actual. Los datos visibles son de demostración.')

# La validacion de carga sigue en pagina aparte: en la anterior no entra la
# figura y quedaria sola al pie de la pagina.
doc.add_page_break()
doc.add_paragraph('Vuelva a Validar hoja de ruta y presione Validar carga, o haga doble clic sobre el renglón pendiente del checklist: RND abre la pantalla donde se resuelve. Tilde cada pallet a medida que lo sube al camión; también puede escribir o escanear el código en el campo Código y presionar Enter, o usar Marcar todos cuando el camión ya está completo.')
doc.add_paragraph('Presione Guardar: los pallets tildados quedan CARGADO y los destildados vuelven a ARMADO. Recién al guardar queda registrado quién cargó cada pallet y cuándo. Cancelar no guarda nada y los pallets conservan su estado.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(2)
p.add_run().add_picture(str(shots/'05_validar_carga.png'), width=Inches(4.35))
p=doc.add_paragraph('Figura 5. Validación de la carga: cada pallet se marca a medida que sube al camión. Los datos visibles son de demostración.', style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER

step('5. Revisar la hoja e imprimir', [
    'En Ver Hoja de Ruta, seleccione la fecha y la ruta y presione Actualizar. Revise el cliente, el comprobante, el producto, las cantidades, los KG, los bultos y los recursos asignados.',
    'Si los recursos están completos y hay pedidos, use Vista previa PDF para revisar la hoja durante la preparación. Cuando la hoja esté en estado LISTA, el botón se muestra como Imprimir PDF. Para marcarla LISTA, entre en Validar hoja de ruta y resuelva los requisitos pendientes del checklist. La impresión y el estado LISTA son controles independientes.',
    'Mientras quede un requisito pendiente, Marcar LISTA y Confirmar salida a reparto permanecen deshabilitados. Para corregir un requisito puntual, haga doble clic sobre la fila: RND abre la pantalla donde se resuelve.',
    'Con la hoja en LISTA, presione Confirmar salida a reparto: el camión salió a reparto y la hoja queda en estado DESPACHADA. Esta acción no marca los pedidos como entregados, y una hoja despachada no retrocede automáticamente.'
], '06_ver_hoja_ruta.png', 'Figura 6. Hoja en estado LISTA con la acción Imprimir PDF habilitada. Los datos visibles son de demostración.')

# Generated output as portrait, plus print guidance.
doc.add_page_break(); doc.add_heading('6. El PDF de reparto', level=1)
doc.add_paragraph('Al presionar Imprimir PDF, RND genera y abre el documento de la fecha y ruta seleccionadas. Desde el visor, revise el contenido y elija la impresora o la opción para guardar una copia, según lo que necesite.')
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(shots/'07_hoja_de_ruta_impresa.png'), width=Inches(4.35))
p=doc.add_paragraph('Figura 7. Ejemplo del PDF de reparto generado por RND. Los datos visibles son de demostración.', style='Caption'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('La impresión requiere pedidos para la fecha y ruta elegidas, además de un chofer y un camión válidos. Si falta algún dato, vuelva a la pantalla de asignación o a la bandeja para corregirlo.')
doc.add_paragraph('Lista para Sergio genera otro PDF, organizado por cliente y destino, con los números de factura agrupados por proveedor. El chofer debe llevar también las facturas originales.')

doc.add_heading('7. Si algo no aparece o no permite continuar', level=1)
for t in [
    'La vista previa no muestra datos: confirme proveedor, archivo y hoja de Excel. Para PDF o imagen, revise la calidad del documento y el mensaje de procesamiento.',
    'La lectura del archivo tiene errores: corrija la vista previa antes de grabar. No grabe filas que no haya podido verificar.',
    'No hay pedidos para asignar: revise la fecha y la ruta; confirme que organizó las facturas en esa ruta.',
    'La hoja no muestra pedidos: vuelva a presionar Actualizar y confirme fecha y ruta.',
    'No se habilita la impresión: confirme que haya pedidos y que la hoja tenga chofer y camión asignados.',
    'No se habilita Marcar LISTA: abra el checklist de Validar hoja de ruta y resuelva los requisitos pendientes.',
    'El checklist pide la carga y no deja marcar LISTA: presione Validar carga desde el checklist y marque cada pallet como CARGADO. Si quedan líneas sin asignar, el doble clic lo lleva a Armado de pallets.',
    'Un pallet no aparece en Armado de pallets: si ya está CARGADO no se muestra ahí. No se perdió: sigue visible en Validar carga y en el checklist.',
    'Confirmar preparación y el pallet sigue pendiente: es lo esperado. Confirmar preparación no cambia el estado; el pallet pasa a CARGADO en Validar carga.',
    'La importación se interrumpió: revise el resumen y los pedidos ya existentes antes de volver a cargar el archivo.'
]: doc.add_paragraph(t, style='List Bullet')

doc.core_properties.title='Manual de usuario - Importar pedidos e imprimir la hoja de ruta'
doc.core_properties.subject='Circuito de reparto en RND'
doc.core_properties.author='Vogel Consultoría'
doc.save(out)
print(out)
