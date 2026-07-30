# coding=utf-8
import datetime
import decimal
import operator
import re

import xlsxwriter
from PyQt5 import QtCore
from PyQt5.QtCore import QAbstractTableModel, Qt, QVariant
from PyQt5.QtGui import QFont, QColor
from PyQt5.QtWidgets import QTableWidget, QTableWidgetItem, QTableView, QApplication, QHeaderView
from openpyxl.reader.excel import load_workbook

from . import Ventanas
from .utiles import EsVerdadero, AbrirArchivo, saveFileDialog, inicializar_y_capturar_excepciones


class Grilla(QTableWidget):
    # columnas a ocultar
    columnasOcultas = []

    # lista con las cabeceras de la grilla
    cabeceras = []

    # tabla desde la cual obtener los datos
    tabla = None

    # campos de la tabla
    campos = None

    # condiciones para filtrar los datos
    condiciones = None

    # cantidad de registros a mostrar
    limite = 100

    # columnas habilitadas
    columnasHabilitadas = []

    # campos tabla
    camposTabla = None

    # valores a cargar
    data = None

    # indica si esta en la grilla
    engrilla = False

    # indica si las columnas son seleccionables o no
    enabled = True

    # widget para las columnas
    widgetCol = {}

    # color para la columna
    backgroundColorCol = {}

    # tamaño de la fuente
    tamanio = 10

    # emit signal
    keyPressed = QtCore.pyqtSignal(int, bool, bool)

    when = QtCore.pyqtSignal()

    # modificadores de teclas
    shift = False
    ctrl = False

    # indica si permite agregar registros
    permiteagregar = True

    # controla la tecla enter para pasar al siguiente campo
    controlaenter = False

    # items nuevo para cuando se agrega una fila
    items_nuevos = []

    LanzarExcepciones = True

    formatos = {}
    
    copy_data = []

    def __init__(self, *args, **kwargs):

        QTableWidget.__init__(self, *args)
        if 'tamanio' in kwargs:
            self.tamanio = kwargs['tamanio']
        font = QFont()
        font.setPointSizeF(self.tamanio)
        self.setFont(font)
        if 'habilitarorden' in kwargs:
            self.setSortingEnabled(kwargs['habilitarorden'])
        else:
            self.setSortingEnabled(True)
        if 'enabled' in kwargs:
            self.setEnabled(kwargs['enabled'])
        self.setAlternatingRowColors(True)
        self.setSelectionBehavior(QTableWidget.SelectRows)
        self.setSelectionMode(QTableWidget.SingleSelection)
        self.setShowGrid(False)
        self.setWordWrap(False)
        self.verticalHeader().setVisible(False)
        self.verticalHeader().setDefaultSectionSize(28)
        self.horizontalHeader().setHighlightSections(False)

    def ArmaCabeceras(self, cabeceras=None):

        if not cabeceras:
            cabeceras = self.cabeceras

        self.setColumnCount(len(cabeceras))

        for col in range(0, len(cabeceras)):
            self.setHorizontalHeaderItem(col, QTableWidgetItem(cabeceras[col]))

        self.cabeceras = cabeceras
        self.OcultaColumnas()
        self.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.horizontalHeader().setStretchLastSection(False)
        self.resizeRowsToContents()
        self.resizeColumnsToContents()

    def AgregaItem(self, items=None,
                   backgroundColor=QColor(255, 255, 255), readonly=False, formatea_miles=False):

        if items:
            col = 0
            cantFilas = self.rowCount() + 1
            self.setRowCount(cantFilas)
            for x in items:
                flags = QtCore.Qt.ItemIsSelectable
                if isinstance(x, bool):
                    item = QTableWidgetItem(x)
                    if x:
                        item.setCheckState(QtCore.Qt.Checked)
                    else:
                        item.setCheckState(QtCore.Qt.Unchecked)
                    self.formatos[col] = 'Bool'
                elif isinstance(x, (int, float, decimal.Decimal)):
                    if formatea_miles:
                        # Formatear con separador de miles y decimales
                        if isinstance(x, int):
                            texto = "{:,}".format(x).replace(",", ".")
                        else:
                            texto = "{:,.2f}".format(float(x)).replace(",", "X").replace(".", ",").replace("X", ".")
                    else:
                        texto = str(x)
                    item = QTableWidgetItem(texto)
                    item.setTextAlignment(Qt.AlignRight)
                    self.formatos[col] = 'Decimal'
                # en caso de que sea formato de fecha
                elif isinstance(x, (datetime.date)):
                    fecha = x.strftime('%d/%m/%Y')
                    item = QTableWidgetItem(fecha)
                    self.formatos[col] = 'Date'
                # en caso de que sea formato de hora
                elif isinstance(x, (datetime.time)):
                    fecha = x.strftime('%H:%M:%S')
                    item = QTableWidgetItem(fecha)
                    self.formatos[col] = 'Time'
                elif isinstance(x, (bytes)):
                    if EsVerdadero(x):
                        item = 'Si'
                    else:
                        item = 'No'
                    item = QTableWidgetItem(QTableWidgetItem(x))
                    self.formatos[col] = 'Bytes'
                else:
                    item = QTableWidgetItem(QTableWidgetItem(x))
                    self.formatos[col] = 'String'

                if readonly:
                    flags = QtCore.Qt.ItemIsSelectable
                elif col in self.columnasHabilitadas:
                    if isinstance(x, (bool)):
                        flags |= QtCore.Qt.ItemIsUserCheckable | QtCore.Qt.ItemIsEditable | QtCore.Qt.ItemIsEnabled
                    else:
                        flags |= QtCore.Qt.ItemIsEnabled | QtCore.Qt.ItemIsEditable
                else:
                    if self.enabled and not readonly:
                        flags |= QtCore.Qt.ItemIsEnabled

                item.setFlags(flags)
                if col in self.backgroundColorCol:
                    item.setBackground(self.backgroundColorCol[col])

                if backgroundColor and isinstance(backgroundColor, QColor):
                    item.setBackground(backgroundColor)

                self.setItem(cantFilas - 1, col, item)
                # if self.widgetCol:
                #     self.ArmaWidgetCol(col)
                if col in self.widgetCol:
                    widgetColumna = self.widgetCol[col]
                    self.setItemDelegateForColumn(col, widgetColumna)
                    # self.setItemDelegate(widgetColumna)
                    # self.setCellWidget(cantFilas - 1, col, widgetColumna)
                col += 1
            if cantFilas <= 25:
                self.resizeRowsToContents()
                self.resizeColumnsToContents()

    # def ArmaWidgetCol(self, col):
    #     for x in range(self.rowCount()):
    #         if col in self.widgetCol:
    #             # self.setCellWidget(x, col, self.widgetCol[col])
    #             self.setItemDelegate(widgetColumna)
    #
    def OcultaColumnas(self):
        for x in self.columnasOcultas:
            if not isinstance(x, int):
                numCol = self.cabeceras.index(x)
            else:
                numCol = x
            self.hideColumn(numCol)

    def ModificaItem(self, valor, fila, col, backgroundColor=None, readonly=False):
        """

        :param fila: la fila que se quiere modificar
        :param valor: valor a modificar
        :type col: entero en caso de indicar un numero de columna y string si quiero el nombre
        """
        if isinstance(col, str):
            col = self.cabeceras.index(col)

        if isinstance(valor, bool):
            item = QTableWidgetItem(valor)
            if valor:
                item.setCheckState(QtCore.Qt.Checked)
            else:
                item.setCheckState(QtCore.Qt.Unchecked)
        elif isinstance(valor, (int, float, decimal.Decimal)):
            item = QTableWidgetItem(str(valor))
        elif isinstance(valor, (datetime.date)):
            fecha = valor.strftime('%d/%m/%Y')
            item = QTableWidgetItem(fecha)
        else:
            item = QTableWidgetItem(valor)

        if not isinstance(col, int):
            numCol = self.cabeceras.index(col)
        else:
            numCol = col
        flags = QtCore.Qt.ItemIsSelectable
        if readonly:
            flags = QtCore.Qt.ItemIsSelectable
        elif col in self.columnasHabilitadas:
            if isinstance(valor, (bool)):
                flags |= QtCore.Qt.ItemIsUserCheckable | QtCore.Qt.ItemIsEditable | QtCore.Qt.ItemIsEnabled
            else:
                flags |= QtCore.Qt.ItemIsEnabled | QtCore.Qt.ItemIsEditable
        else:
            if self.enabled and not readonly:
                flags |= QtCore.Qt.ItemIsEnabled

        item.setFlags(flags)

        if col in self.backgroundColorCol:
            item.setBackground(self.backgroundColorCol[col])
        else:
            # si no indico un color coloco el color que tiene el item
            if self.itemAt(fila, col):
                item.setBackground(self.itemAt(fila, col).background().color())

        self.setItem(fila, numCol, item)
        if self.rowCount() <= 25:
            self.resizeColumnsToContents()

    def ObtenerItem(self, fila, col):

        if isinstance(col, int):
            numCol = col
        else:
            numCol = self.cabeceras.index(col)

        try:
            item = self.item(fila, numCol)
            if item.checkState() == QtCore.Qt.Checked:
                item = True
            else:
                item = item.text()
                item = item.replace(',', '.')  # if item else 0

        except:
            item = ''

        # return item.replace(',','.') if item else 0
        return item

    def ObtenerItemNumerico(self, fila, col, configuracion_numero=False):
        if configuracion_numero:
            import locale
            import os

            try:
                # Intenta usar el locale por defecto del sistema
                locale.setlocale(locale.LC_NUMERIC, '')
            except locale.Error:
                # En caso de error (Windows sin locale), usa uno por defecto o el del sistema
                lang, _ = locale.getdefaultlocale()
                try:
                    locale.setlocale(locale.LC_NUMERIC, lang)
                except locale.Error:
                    print("No se pudo cargar el locale del sistema, usando 'C'")
                    locale.setlocale(locale.LC_NUMERIC, 'C.UTF-8')  # Fallback seguro
        
        if isinstance(col, int):
            numCol = col
        else:
            numCol = self.cabeceras.index(col)

        try:
            item = self.item(fila, numCol)
            if item:
                if item.checkState() == QtCore.Qt.ItemIsUserCheckable:
                    item = True
                else:
                    item = item.text()
                    item = re.sub(r"[^-0123456789\.,]", "", item)
                
                if configuracion_numero:
                    item = locale.atof(item)
                else:
                    item = float(item)
            else:
                item = 0
        except Exception as e:
            pass
            # Ventanas.showAlert("INFO", f"Error al convertir a numero {e} {col} {fila}")

        # return item.replace(',','.') if item else 0
        return item

    @inicializar_y_capturar_excepciones
    def ObtenerItemFecha(self, fila, col):
        if isinstance(col, int):
            numCol = col
        else:
            numCol = self.cabeceras.index(col)

        try:
            item = self.item(fila, numCol).text().replace('-', '/').strip()
            item = datetime.datetime.strptime(item, '%d/%m/%Y')
        except:
            item = ''

        return item

    def CargaDatos(self, avance=None):
        self.blockSignals(True)
        self.setRowCount(0)
        self.blockSignals(False)

    def focusInEvent(self, *args, **kwargs):
        self.engrilla = True
        if self.rowCount() == 0:
            self.AgregaNuevo()
        super().focusInEvent(*args, **kwargs)
        self.when.emit()

    def focusOutEvent(self, *args, **kwargs):
        self.engrilla = False
        super().focusOutEvent(*args, **kwargs)

    def HabilitaColumnas(self, fila=1):

        for col in range(self.columnCount()):
            valor = self.ObtenerItem(fila=fila, col=col)
            self.ModificaItem(valor=valor, fila=fila, col=col)

    def ExportaExcel(self, columnas=None, archivo="", titulo="", nuevo=True, hoja='', fila=0, col=0, avance=None):

        if not columnas:
            columnas = self.cabeceras

        if nuevo:
            archivo = archivo.replace('.', '').replace('/', '')
        if not archivo.startswith("excel"):
            archivo = "excel/" + archivo

        if nuevo:
            cArchivo = saveFileDialog(filename=archivo,
                                      files="Archivos de Excel (*.xlsx)")
        else:
            cArchivo = archivo
        # cArchivo = QFileDialog.getSaveFileName(caption="Guardar archivo", directory="", filter="*.XLSX")
        if not cArchivo:
            return

        if nuevo:
            workbook = xlsxwriter.Workbook(cArchivo)
        else:
            try:
                # Carga el archivo
                workbook = load_workbook(cArchivo)
            except:
                Ventanas.showAlert("Sistema", f"Ocurrio un error al intentar abrir el archivo {cArchivo}")
                return

        if hoja:
            # Selecciona la hoja por su nombre
            worksheet = workbook[hoja]
        else:
            #creamos una hoja nueva
            worksheet = workbook.add_worksheet()

        formato_fecha = workbook.add_format({'num_format': 'dd/mm/yyyy'})
        # fila = 0
        # col = 0
        if titulo:
            worksheet.write(fila, col, titulo)
            fila += 2

        for c in columnas:
            worksheet.write(fila, col, c)
            col += 1

        fila += 1
        total_rows = self.rowCount()
        for row in range(total_rows):
            col = 0
            for c in columnas:
                dato = self.ObtenerItem(fila=row, col=c)
                if isinstance(dato, bool):
                    dato = 'SI' if dato else 'NO'
                else:
                    dato = dato.strip()
                try:
                    dato = float(dato)
                except:
                    if dato.isdigit():
                        dato = int(dato)
                if self.formatos[col] == 'Date':
                    try:
                        worksheet.write_datetime(fila, col, datetime.datetime.strptime(dato.replace("-", "/"), '%d/%m/%Y').date(),
                                                formato_fecha)
                    except:
                        worksheet.write(fila, col, dato, formato_fecha)
                else:
                    worksheet.write(fila, col, dato)
                col += 1
            fila += 1
            # actualizar barra de progreso si se provee
            try:
                if avance is not None and total_rows > 0:
                    porcentaje = int((row + 1) / total_rows * 100)
                    # avance puede ser un widget Avance con método actualizar
                    if hasattr(avance, 'actualizar'):
                        avance.actualizar(porcentaje)
                    # asegurarse de procesar eventos para mantener UI responsive
                    QApplication.processEvents()
            except Exception:
                pass
        # cabeceras_excel = [{'header': x} for x in columnas]
        # worksheet.add_table(0, 0, fila, col-1, cabeceras_excel)
        workbook.close()
        AbrirArchivo(cArchivo)

    def obtiene_datos_lista(self, campos=None):
        datos_retorno = []
        if campos is None:
            campos = []
        for row in range(self.rowCount()):
            for c in campos:
                dato = self.ObtenerItem(fila=row, col=c)
                if isinstance(dato, bool):
                    dato = 'SI' if dato else 'NO'
                else:
                    dato = dato.strip()
                try:
                    dato = float(dato)
                except:
                    if dato.isdigit():
                        dato = int(dato)
                datos_retorno.append(dato)
        return [x for x in datos_retorno]

    def keyPressEvent(self, event):
        super(Grilla, self).keyPressEvent(event)
        self.shift = False
        self.ctrl = False
        if (event.modifiers() & Qt.ShiftModifier):
            self.shift = True
        if (event.modifiers() & Qt.ControlModifier):
            self.ctrl = True

        col = self.currentColumn()
        row = self.currentRow()
        if event.key() in [Qt.Key_Return, Qt.Key_Enter] and self.controlaenter:
            if col + 1 == self.columnCount() and row + 1 == self.rowCount():
                if self.permiteagregar:
                    if self.items_nuevos:
                        item = self.items_nuevos
                    else:
                        item = ['' for x in self.cabeceras]
                    self.AgregaItem(item)
                    self.Activar(row=self.rowCount(), col=0)
            elif col == self.columnCount():
                self.Activar(row=row + 1, col=0)
            else:
                self.Activar(col=col + 1)
        elif event.key() == Qt.Key_Down and row + 1 == self.rowCount():
            if self.permiteagregar:
                if self.items_nuevos:
                    item = self.items_nuevos
                else:
                    item = ['' for x in self.cabeceras]
                self.AgregaItem(item)
                # self.Activar(row=row + 1, col=0)
        elif event.key() == Qt.Key_C and self.ctrl:
            self.copySelection()
        elif event.key() == Qt.Key_V and self.ctrl:
            self.pasteSelection()
        self.keyPressed.emit(event.key(), self.shift, self.ctrl)

    def SumaTodo(self, col=None):
        total = 0.
        for row in range(self.rowCount()):
            monto = float(self.ObtenerItem(fila=row, col=col) or '0')
            total += monto
        return total

    def Activar(self, row=-1, col=0):
        if row == -1:
            row = self.rowCount() - 1
        self.setCurrentCell(row, col)
        self.setFocus()

    def BorraItemCodigo(self, codigo=None, col=0):
        if not codigo:
            return 0

        cant = 0
        for row in range(self.rowCount()):
            dato = self.ObtenerItem(fila=row, col=col)
            if dato == codigo:
                self.removeRow(row)
                cant += 1
        return cant

    def AgregaNuevo(self):
        items = ['' for x in self.cabeceras]
        self.AgregaItem(items)

    def copySelection(self):
        selection = self.selectedIndexes()
        if selection:
            rows = sorted(index.row() for index in selection)
            columns = sorted(index.column() for index in selection)
            rowcount = rows[-1] - rows[0] + 1
            colcount = columns[-1] - columns[0] + 1
            table = [[''] * colcount for _ in range(rowcount)]
            for index in selection:
                row = index.row() - rows[0]
                column = index.column() - columns[0]
                table[row][column] = index.data()
            self.copy_data = table
        return

    def pasteSelection(self):
        selection = self.selectedIndexes()
        if selection:
            model = self.model()
            rows = sorted(index.row() for index in selection)
            columns = sorted(index.column() for index in selection)
            if len(rows) == 1 and len(columns) == 1:
                fila = rows[0]
                columna = columns[0]
                self.setCurrentCell(0, 0)
                for f in self.copy_data:
                    for c in f:
                        self.ModificaItem(valor=c, fila=fila, col=columna)
                        columna += 1
                        self.setCurrentCell(fila, columna)
                    fila += 1
        return
    
    def limpiarGrilla(self):
        self.setSortingEnabled(False)
        self.setRowCount(0)
        
    def filaSeleccionada(self):
        try:
            return self.currentRow()
        except:
            return 0

class MiTableModel(QAbstractTableModel):
    # columnas habilitadas
    columnasHabilitadas = []
    # _colorColumn = {2:QColor(255, 255, 204),4:QColor(255, 128, 128)}
    _colorColumn = {}

    def __init__(self, datain, headerdata, parent=None, *args):
        QAbstractTableModel.__init__(self, parent, *args)
        self.arraydata = datain
        self.headerdata = headerdata

    def rowCount(self, parent=None, *args, **kwargs):
        return len(self.arraydata)

    def columnCount(self, parent=None, *args, **kwargs):
        if self.arraydata:
            return len(self.arraydata[0])
        else:
            return 0

    def data(self, index, role=None):
        if not index.isValid():
            return None
        value = self.arraydata[index.row()][index.column()]
        if role == Qt.EditRole:
            return value
        elif role == Qt.DisplayRole:
            return value
        elif role == Qt.ItemIsSelectable:
            return value
        elif role == Qt.BackgroundRole and index.column() in (self._colorColumn):
            return self._colorColumn[index.column()]
        return QVariant()

    def headerData(self, col, orientation, role=None):
        if orientation == Qt.Horizontal and role == Qt.DisplayRole:
            return QVariant(self.headerdata[col])
        return QVariant()

    def sort(self, col, order=None):
        self.layoutAboutToBeChanged.emit()
        self.arraydata = sorted(self.arraydata, key=operator.itemgetter(col))
        if order == Qt.DescendingOrder:
            self.arraydata.reverse()
        self.layoutChanged.emit()

    def flags(self, index):
        flags = super(self.__class__, self).flags(index)
        flags |= Qt.ItemIsSelectable | QtCore.Qt.ItemIsEnabled
        if index.column() in self.columnasHabilitadas:
            flags |= Qt.ItemIsEditable
        return flags

    def setData(self, index, value, role=None):
        if role == Qt.EditRole:
            row = index.row()
            column = index.column()
            ch = (value)
            self.arraydata[row][column] = ch
            self.dataChanged.emit(index, index)
            return True

    def ModificaDato(self, fila=0, col=0, valor=None):
        self.arraydata[fila][col] = valor


class MiQTableView(QTableView):
    def __init__(self, *args, **kwargs):
        QTableView.__init__(self, *args, **kwargs)  # Use QTableView constructor


class GrillaModel(MiQTableView):
    # columnas a ocultar
    columnasOcultas = []

    # emit signal
    keyPressed = QtCore.pyqtSignal(int)

    def __init__(self, parent=None):
        super().__init__()
        font = QFont()
        font.setPointSizeF(12)
        self.setFont(font)

    def ObtenerItem(self, fila=0, col=0):
        # en caso de que le pase el nombre de la columnna busco el indice del header
        if isinstance(col, (str,)):
            col = self.model().headerdata.index(col)

        index = self.model().index(fila, col)

        return index.data()

    def ModificaItem(self, fila=0, col=0, valor=''):
        # en caso de que le pase el nombre de la columnna busco el indice del header
        if isinstance(col, (str,)):
            col = self.model().headerdata.index(col)

        index = self.model().index(fila, col)
        if isinstance(valor, (int, float, decimal.Decimal)):
            valor = str(valor)

        self.model().setData(index, valor, role=Qt.EditRole)

    def OcultarColumnas(self):
        for i in self.columnasOcultas:
            if isinstance(i, (str,)):
                col = self.model().headerdata.index(i)
            else:
                col = i

            self.hideColumn(col)

    def keyPressEvent(self, event):
        super(GrillaModel, self).keyPressEvent(event)
        self.keyPressed.emit(event.key())
